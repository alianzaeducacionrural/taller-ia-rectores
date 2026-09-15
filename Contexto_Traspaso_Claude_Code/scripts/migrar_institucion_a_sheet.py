"""
Script validado para migrar un Excel de origen (Info_Comite/COR_<n>_<NOMBRE>.xlsx)
a un workbook .xlsx listo para subir/convertir a Google Sheets, agregando la
pestaña calculada "Competencias Débiles".

Uso:
    python3 migrar_institucion_a_sheet.py <ruta_excel_origen.xlsx> <ruta_salida.xlsx> [--caldas]

Ejemplo (instituciones normales, incluye Detalle Estudiantes):
    python3 migrar_institucion_a_sheet.py Info_Comite/COR_1_GIOVANNI_MONTINI.xlsx "Giovanni Montini.xlsx"

Ejemplo (Caldas / consolidado — NUNCA incluir Detalle Estudiantes, porque en el
archivo fuente esa pestaña trae el detalle de TODAS las instituciones):
    python3 migrar_institucion_a_sheet.py Info_Comite/COR_100_CALDAS.xlsx "Caldas (consolidado).xlsx" --caldas

Requiere: pip install openpyxl
"""
import sys
import argparse
import openpyxl
from openpyxl import Workbook
from collections import defaultdict


def load_sheet_rows(path, sheet):
    wb = openpyxl.load_workbook(path, data_only=True)
    return list(wb[sheet].iter_rows(values_only=True))


def competencias_debiles_rows(path, top_n=3):
    """Agrupa 'Estadísticas Preguntas' por (materia, competencia), promedia el
    % de acierto, y devuelve las top_n competencias con menor acierto promedio."""
    rows = load_sheet_rows(path, 'Estadísticas Preguntas')
    header = rows[0]
    idx = {h: i for i, h in enumerate(header)}
    agg = defaultdict(list)
    meta = {}
    for r in rows[1:]:
        comp = r[idx['competencia']]
        pct = r[idx['% acierto']]
        materia_code = r[idx['# Pregunta']]  # ej: MT1, LC3...
        materia = ''.join([c for c in materia_code if not c.isdigit()])
        agg[(materia, comp)].append(pct)
        if (materia, comp) not in meta:
            meta[(materia, comp)] = {
                'afirmacion': r[idx['afirmacion']],
                'evidencia': r[idx['evidencia']],
            }
    avgs = []
    for k, vals in agg.items():
        avg_pct = sum(vals) / len(vals)
        avgs.append((k[0], k[1], round(avg_pct * 100, 1), meta[k]['afirmacion'], meta[k]['evidencia']))
    avgs.sort(key=lambda x: x[2])
    out = [('materia', 'competencia', '% acierto promedio', 'afirmacion', 'evidencia')]
    out.extend(avgs[:top_n])
    return out


def build_institution_workbook(src_path, out_path, include_detalle=True):
    wb_out = Workbook()
    wb_out.remove(wb_out.active)

    ws = wb_out.create_sheet('Resultados por Materia')
    for row in load_sheet_rows(src_path, 'Resultados por Materia'):
        ws.append(row)

    ws = wb_out.create_sheet('NV Desempeño')
    for row in load_sheet_rows(src_path, 'NV DESEMPEÑO'):
        ws.append(row)

    ws = wb_out.create_sheet('Estadísticas Preguntas')
    for row in load_sheet_rows(src_path, 'Estadísticas Preguntas'):
        ws.append(row)

    ws = wb_out.create_sheet('Competencias Débiles')
    for row in competencias_debiles_rows(src_path):
        ws.append(row)

    if include_detalle:
        ws = wb_out.create_sheet('Detalle Estudiantes')
        for row in load_sheet_rows(src_path, 'Detalle Estudiantes (ref)'):
            ws.append(row)

    wb_out.save(out_path)
    print('OK ->', out_path)


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('origen')
    parser.add_argument('salida')
    parser.add_argument('--caldas', action='store_true',
                         help='Usar para el archivo consolidado de Caldas: excluye la pestaña Detalle Estudiantes')
    args = parser.parse_args()
    build_institution_workbook(args.origen, args.salida, include_detalle=not args.caldas)
